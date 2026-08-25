"""Convert supported document formats to PDF for pipeline ingestion."""

from __future__ import annotations

import csv
import platform
import shutil
from pathlib import Path
from typing import Union


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".pptx",
    ".ppt",
    ".xlsx",
    ".xls",
    ".txt",
    ".csv",
}

CONTENT_TYPE_MAP = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".ppt": "application/vnd.ms-powerpoint",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".txt": "text/plain",
    ".csv": "text/csv",
}


def is_supported_document(path: Union[str, Path]) -> bool:
    """Return True if the file extension is a supported document type."""
    return Path(path).suffix.lower() in SUPPORTED_EXTENSIONS


def get_content_type(filename: str) -> str:
    """Return the MIME content type for a supported document."""
    ext = Path(filename).suffix.lower()
    return CONTENT_TYPE_MAP.get(ext, "application/octet-stream")


def convert_to_pdf(input_path: Path, output_path: Path) -> Path:
    """
    Convert a supported document to PDF.

    For PDFs the file is copied. For other supported types the text
    is extracted and written into a simple PDF so the existing
    pdfplumber / PaddleOCR pipeline can operate on it unchanged.

    Returns the path to the output PDF.
    """
    ext = input_path.suffix.lower()

    if ext == ".pdf":
        shutil.copy2(input_path, output_path)
        return output_path

    text = _extract_text(input_path, ext)
    _write_text_pdf(text, output_path)
    return output_path


# ============================================================================
# Text extraction
# ============================================================================


def _extract_text(input_path: Path, ext: str) -> str:
    if ext in (".txt", ".csv"):
        return _read_text_file(input_path)
    elif ext == ".docx":
        return _extract_docx_text(input_path)
    elif ext == ".doc":
        return _extract_doc_text(input_path)
    elif ext == ".pptx":
        return _extract_pptx_text(input_path)
    elif ext == ".ppt":
        return _extract_ppt_text(input_path)
    elif ext == ".xlsx":
        return _extract_xlsx_text(input_path)
    elif ext == ".xls":
        return _extract_xls_text(input_path)
    else:
        raise ValueError(
            f"Unsupported extension for text extraction: {ext}"
        )


def _read_text_file(input_path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return input_path.read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return input_path.read_text(encoding="utf-8", errors="replace")


def _extract_docx_text(input_path: Path) -> str:
    from docx import Document

    doc = Document(str(input_path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(
                " | ".join(cell.text for cell in row.cells)
            )
    return "\n".join(parts)


def _extract_doc_text(input_path: Path) -> str:
    try:
        import textract
        return textract.process(
            str(input_path)
        ).decode("utf-8", errors="replace")
    except ImportError as exc:
        raise ImportError(
            "textract is required for .doc files. "
            "Install it or convert to .docx first."
        ) from exc


def _extract_pptx_text(input_path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(input_path))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
        parts.append("")
    return "\n".join(parts)


def _extract_ppt_text(input_path: Path) -> str:
    try:
        import textract
        return textract.process(
            str(input_path)
        ).decode("utf-8", errors="replace")
    except ImportError as exc:
        raise ImportError(
            "textract is required for .ppt files. "
            "Install it or convert to .pptx first."
        ) from exc


def _extract_xlsx_text(input_path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(
        str(input_path),
        read_only=True,
        data_only=True,
    )
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            values = [
                str(cell) if cell is not None else ""
                for cell in row
            ]
            parts.append(" | ".join(values))
    return "\n".join(parts)


def _extract_xls_text(input_path: Path) -> str:
    try:
        import xlrd

        wb = xlrd.open_workbook(str(input_path))
        parts: list[str] = []
        for sheet in wb.sheets():
            for row in range(sheet.nrows):
                values = [
                    str(sheet.cell_value(row, col))
                    for col in range(sheet.ncols)
                ]
                parts.append(" | ".join(values))
        return "\n".join(parts)
    except ImportError as exc:
        raise ImportError(
            "xlrd is required for .xls files. "
            "Install it or convert to .xlsx first."
        ) from exc


# ============================================================================
# PDF generation
# ============================================================================


def _write_text_pdf(text: str, output_path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    font_path = _find_system_font()
    if font_path:
        try:
            pdf.add_font(
                "DocumentFont",
                "",
                str(font_path),
                uni=True,
            )
            pdf.set_font("DocumentFont", size=12)
        except Exception:
            pdf.set_font("Helvetica", size=12)
    else:
        pdf.set_font("Helvetica", size=12)

    pdf.multi_cell(0, 5, text)
    pdf.output(str(output_path))


def _find_system_font() -> Union[Path, None]:
    system = platform.system()
    candidates: list[Path] = []

    if system == "Windows":
        candidates = [
            Path("C:\\Windows\\Fonts\\arial.ttf"),
            Path("C:\\Windows\\Fonts\\calibri.ttf"),
            Path("C:\\Windows\\Fonts\\segoeui.ttf"),
        ]
    elif system == "Darwin":
        candidates = [
            Path("/Library/Fonts/Arial.ttf"),
            Path("/System/Library/Fonts/Helvetica.ttc"),
        ]
    else:
        candidates = [
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
        ]

    for path in candidates:
        if path.exists():
            return path
    return None
